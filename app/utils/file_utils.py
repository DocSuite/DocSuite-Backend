from pathlib import Path
from dataclasses import dataclass, field
import mimetypes

from app.core.exceptions import DocSuiteException


@dataclass
class DocumentElement:
    type: str
    content: str
    page: int | None = None
    section: str | None = None
    metadata: dict[str, str | int | None] = field(default_factory=dict)
    image_bytes: bytes | None = None
    image_mime_type: str | None = None


@dataclass
class DocumentExtraction:
    text: str
    elements: list[DocumentElement]


def extract_text_from_document(file_path: Path) -> str:
    return extract_structured_document(file_path).text


def extract_structured_document(file_path: Path) -> DocumentExtraction:
    suffix = file_path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return _extract_text_file(file_path)
    if suffix == ".pdf":
        return _extract_pdf(file_path)
    if suffix == ".docx":
        return _extract_docx(file_path)
    raise DocSuiteException("Formato no soportado. Usa PDF, DOCX, TXT o MD.")


def _format_elements(elements: list[DocumentElement]) -> str:
    blocks: list[str] = []
    for index, element in enumerate(elements, start=1):
        source_parts = [f"elemento {index}", f"tipo {element.type}"]
        if element.page is not None:
            source_parts.append(f"pagina {element.page}")
        if element.section:
            source_parts.append(f"seccion {element.section}")
        for key, value in element.metadata.items():
            if value is not None:
                source_parts.append(f"{key} {value}")
        blocks.append(f"[Fuente: {' | '.join(source_parts)}]\n{element.content.strip()}")
    return "\n\n".join(blocks).strip()


def _extract_text_file(file_path: Path) -> DocumentExtraction:
    content = file_path.read_text(encoding="utf-8").strip()
    if not content:
        raise DocSuiteException("El documento no contiene texto para analizar.")
    element = DocumentElement(type="text", content=content, section=file_path.suffix.lower().lstrip("."))
    return DocumentExtraction(text=_format_elements([element]), elements=[element])


def _extract_pdf(file_path: Path) -> DocumentExtraction:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DocSuiteException("pypdf no esta instalado", status_code=503) from exc

    reader = PdfReader(str(file_path))
    elements: list[DocumentElement] = []
    for page_index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            elements.append(DocumentElement(type="text", content=text, page=page_index))

        images = getattr(page, "images", []) or []
        for image_index, image in enumerate(images, start=1):
            image_name = getattr(image, "name", None) or f"imagen_{image_index}"
            image_bytes = getattr(image, "data", None)
            elements.append(
                DocumentElement(
                    type="image",
                    content=(
                        "Imagen o grafico detectado en el documento. "
                        "No se interpreta visualmente en esta version; se registra como evidencia visual pendiente de OCR/vision."
                    ),
                    page=page_index,
                    metadata={"nombre": image_name, "indice": image_index},
                    image_bytes=image_bytes if isinstance(image_bytes, bytes) else None,
                    image_mime_type=_guess_mime_type(image_name),
                )
            )

    if not elements:
        raise DocSuiteException(
            "No se encontro texto ni imagenes extraibles en el PDF. Si es un documento escaneado, se requiere OCR."
        )

    return DocumentExtraction(text=_format_elements(elements), elements=elements)


def _extract_docx(file_path: Path) -> DocumentExtraction:
    try:
        from docx import Document
    except ImportError as exc:
        raise DocSuiteException("python-docx no esta instalado", status_code=503) from exc

    document = Document(str(file_path))
    elements: list[DocumentElement] = []
    current_section: str | None = None

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = paragraph.style.name if paragraph.style is not None else ""
        if style_name.lower().startswith("heading"):
            current_section = text
            elements.append(DocumentElement(type="heading", content=text, section=current_section))
            continue
        elements.append(DocumentElement(type="text", content=text, section=current_section))

    for table_index, table in enumerate(document.tables, start=1):
        rows: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            rows.append(" | ".join(cells))
        if rows:
            elements.append(
                DocumentElement(
                    type="table",
                    content="\n".join(rows),
                    section=current_section,
                    metadata={"tabla": table_index},
                )
            )

    image_parts = [
        rel.target_part
        for rel in document.part.rels.values()
        if "image" in rel.reltype and hasattr(rel.target_part, "blob")
    ]
    for image_index, image_part in enumerate(image_parts, start=1):
        image_name = Path(getattr(image_part, "partname", f"imagen_{image_index}")).name
        elements.append(
            DocumentElement(
                type="image",
                content=(
                    "Imagen o grafico detectado en el documento. "
                    "No se interpreta visualmente en esta version; se registra como evidencia visual pendiente de OCR/vision."
                ),
                section=current_section,
                metadata={"nombre": image_name, "indice": image_index},
                image_bytes=image_part.blob,
                image_mime_type=getattr(image_part, "content_type", None) or _guess_mime_type(image_name),
            )
        )

    if not elements:
        raise DocSuiteException("El DOCX no contiene texto, tablas ni imagenes extraibles.")

    return DocumentExtraction(text=_format_elements(elements), elements=elements)


def _guess_mime_type(filename: str) -> str:
    mime_type, _encoding = mimetypes.guess_type(filename)
    return mime_type or "image/png"
