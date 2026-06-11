import re
import json
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Inches

_PRIMARY = RGBColor(0x1F, 0x49, 0x7D)
_GRAY = RGBColor(0x70, 0x70, 0x70)
_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
_HEADER_BG = "1F497D"
_ALT_ROW_BG = "EBF0FA"
_LIGHT_BG = "F4F7FB"
_SOFT_BLUE = "D9EAF7"

_SEP_ROW_RE = re.compile(r"^\|[\s\-:|]+\|$")
_NUMBERED_RE = re.compile(r"^\d+\.\s+")
_INLINE_RE = re.compile(r"\*\*(.+?)\*\*|\*(.+?)\*|([^*]+)")


def _parse_row(row: str) -> list[str]:
    parts = row.split("|")
    return [c.strip() for c in parts[1:-1]]


def _set_cell_bg(cell, hex_color: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def _set_cell_text(cell, text: str, bold: bool = False, color: RGBColor | None = None, size: int = 9) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.paragraphs[0].clear()
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color


def _add_run_formatted(para, text: str) -> None:
    for m in _INLINE_RE.finditer(text):
        if m.group(1):
            run = para.add_run(m.group(1))
            run.bold = True
        elif m.group(2):
            run = para.add_run(m.group(2))
            run.italic = True
        elif m.group(3):
            para.add_run(m.group(3))


def _add_hr(doc: Document) -> None:
    para = doc.add_paragraph()
    p_pr = para._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F497D")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _render_table(doc: Document, raw_lines: list[str]) -> None:
    header: list[str] | None = None
    data_rows: list[list[str]] = []

    for line in raw_lines:
        if _SEP_ROW_RE.match(line):
            continue
        cells = _parse_row(line)
        if header is None:
            header = cells
        else:
            data_rows.append(cells)

    if not header:
        return

    col_count = len(header)
    table = doc.add_table(rows=1, cols=col_count)
    table.style = "Table Grid"

    hdr_cells = table.rows[0].cells
    for j, text in enumerate(header):
        cell = hdr_cells[j]
        cell.paragraphs[0].clear()
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.color.rgb = _WHITE
        run.font.size = Pt(9)
        _set_cell_bg(cell, _HEADER_BG)

    for row_idx, row_data in enumerate(data_rows):
        row = table.add_row()
        for j in range(col_count):
            cell = row.cells[j]
            text = row_data[j] if j < len(row_data) else ""
            cell.paragraphs[0].clear()
            _add_run_formatted(cell.paragraphs[0], text)
            cell.paragraphs[0].runs[0].font.size = Pt(9) if cell.paragraphs[0].runs else None
            if row_idx % 2 == 1:
                _set_cell_bg(cell, _ALT_ROW_BG)

    doc.add_paragraph()


def markdown_to_docx(content: str, filename: str = "acta") -> BytesIO:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.5)

    lines = content.split("\n")
    i = 0

    while i < len(lines):
        raw = lines[i]
        s = raw.strip()

        if not s:
            i += 1
            continue

        if s.startswith("### "):
            h = doc.add_heading(s[4:].strip(), level=3)
            for run in h.runs:
                run.font.color.rgb = _PRIMARY
            i += 1

        elif s.startswith("## "):
            h = doc.add_heading(s[3:].strip(), level=2)
            for run in h.runs:
                run.font.color.rgb = _PRIMARY
            i += 1

        elif s.startswith("# "):
            h = doc.add_heading(s[2:].strip(), level=1)
            for run in h.runs:
                run.font.color.rgb = _PRIMARY
            i += 1

        elif s == "---":
            _add_hr(doc)
            i += 1

        elif s.startswith("|"):
            block: list[str] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            _render_table(doc, block)

        elif s.startswith("- [ ]") or s.startswith("- [x]") or s.startswith("- [X]"):
            checked = s[3].lower() == "x"
            text = s[6:].strip()
            symbol = "☑" if checked else "☐"
            para = doc.add_paragraph(style="List Bullet")
            run = para.add_run(f"{symbol}  {text}")
            run.font.size = Pt(10)
            i += 1

        elif s.startswith("- "):
            para = doc.add_paragraph(style="List Bullet")
            _add_run_formatted(para, s[2:].strip())
            i += 1

        elif _NUMBERED_RE.match(s):
            text = _NUMBERED_RE.sub("", s)
            para = doc.add_paragraph(style="List Number")
            _add_run_formatted(para, text)
            i += 1

        elif s.startswith("*") and s.endswith("*") and not s.startswith("**"):
            para = doc.add_paragraph()
            run = para.add_run(s[1:-1])
            run.italic = True
            run.font.color.rgb = _GRAY
            run.font.size = Pt(9)
            i += 1

        else:
            para = doc.add_paragraph()
            _add_run_formatted(para, s)
            i += 1

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def analyzer_to_docx(filename: str, mode: str, result_json: str, extracted_text: str) -> BytesIO:
    doc = Document()
    _setup_document(doc)
    _setup_styles(doc)

    title = _safe_result_title(result_json, filename)
    result = _parse_result_json(result_json)
    _render_cover(doc, filename, mode, title, result)

    if mode == "academic":
        _render_academic_analysis(doc, result)
    else:
        _render_general_analysis(doc, result)

    doc.add_page_break()
    doc.add_heading("Extraccion con fuentes", level=2)
    for block in extracted_text.split("\n\n"):
        s = block.strip()
        if not s:
            continue
        if s.startswith("[Fuente:"):
            source, _, content = s.partition("]\n")
            p = doc.add_paragraph()
            run = p.add_run(source + "]")
            run.bold = True
            run.font.color.rgb = _GRAY
            run.font.size = Pt(9)
            if content.strip():
                doc.add_paragraph(content.strip())
        else:
            doc.add_paragraph(s)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _setup_document(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.5)


def _setup_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.color.rgb = _PRIMARY
        style.font.bold = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)


def _render_cover(doc: Document, filename: str, mode: str, title: str, result: dict[str, Any]) -> None:
    label = doc.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = label.add_run("DocAnalyzer")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = _PRIMARY

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Reporte estructurado de analisis documental")
    subtitle_run.font.size = Pt(11)
    subtitle_run.font.color.rgb = _GRAY

    _add_hr(doc)

    title_paragraph = doc.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_paragraph.add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(16)

    doc.add_paragraph()
    meta_rows = [
        ("Archivo", filename),
        ("Modo", "Analizador academico" if mode == "academic" else "Extractor general"),
        ("Idioma detectado", _stringify(result.get("language") or "No detectado")),
        ("Tipo de documento", _stringify(result.get("document_type") or (result.get("identification") or {}).get("document_type") or "No detectado")),
    ]
    _add_key_value_table(doc, meta_rows)
    doc.add_page_break()


def _add_key_value_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for index, (label, value) in enumerate(rows):
        row = table.add_row()
        _set_cell_text(row.cells[0], label, bold=True, color=_PRIMARY)
        _set_cell_text(row.cells[1], value)
        _set_cell_bg(row.cells[0], _SOFT_BLUE)
        if index % 2 == 1:
            _set_cell_bg(row.cells[1], _LIGHT_BG)

    doc.add_paragraph()


def _add_callout(doc: Document, title: str, body: Any) -> None:
    if body in (None, "", []):
        return
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    _set_cell_bg(cell, _LIGHT_BG)
    cell.paragraphs[0].clear()
    title_run = cell.paragraphs[0].add_run(title)
    title_run.bold = True
    title_run.font.color.rgb = _PRIMARY
    body_p = cell.add_paragraph()
    body_p.add_run(_stringify(body))
    doc.add_paragraph()


def _parse_result_json(result_json: str) -> dict[str, Any]:
    try:
        parsed = json.loads(result_json)
    except json.JSONDecodeError:
        return {"executive_summary": result_json}
    return parsed if isinstance(parsed, dict) else {"executive_summary": result_json}


def _safe_result_title(result_json: str, filename: str) -> str:
    result = _parse_result_json(result_json)
    identification = result.get("identification")
    if isinstance(identification, dict) and identification.get("title"):
        return str(identification["title"])
    return str(result.get("title") or filename)


def _add_label_value(doc: Document, label: str, value: Any) -> None:
    if value in (None, "", []):
        return
    p = doc.add_paragraph()
    p.add_run(f"{label}: ").bold = True
    p.add_run(_stringify(value))


def _add_bullets(doc: Document, title: str, items: Any) -> None:
    if not isinstance(items, list) or not items:
        return
    doc.add_heading(title, level=3)
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(_stringify(item))


def _stringify(value: Any) -> str:
    if isinstance(value, str):
        return value
    if isinstance(value, list):
        return "; ".join(_stringify(item) for item in value)
    if isinstance(value, dict):
        return "; ".join(f"{key}: {_stringify(item)}" for key, item in value.items() if item not in (None, "", []))
    return str(value)


def _render_general_analysis(doc: Document, result: dict[str, Any]) -> None:
    doc.add_heading("Resumen", level=2)
    _add_callout(doc, "Resumen ejecutivo", result.get("executive_summary"))
    _add_bullets(doc, "Puntos clave", result.get("key_points"))

    sections = result.get("main_sections")
    if isinstance(sections, list) and sections:
        doc.add_heading("Secciones principales", level=3)
        for section in sections:
            if isinstance(section, dict):
                _add_label_value(doc, str(section.get("title") or "Seccion"), section.get("summary"))

    _render_document_elements(doc, result)
    _add_bullets(doc, "Conclusiones", result.get("conclusions"))
    _add_bullets(doc, "Palabras clave", result.get("keywords"))


def _render_academic_analysis(doc: Document, result: dict[str, Any]) -> None:
    identification = result.get("identification") if isinstance(result.get("identification"), dict) else {}
    research = result.get("research") if isinstance(result.get("research"), dict) else {}
    methodology = result.get("methodology") if isinstance(result.get("methodology"), dict) else {}
    conclusions = result.get("conclusions") if isinstance(result.get("conclusions"), dict) else {}
    thesis = result.get("thesis_relevance") if isinstance(result.get("thesis_relevance"), dict) else {}

    doc.add_heading("Resumen Academico", level=2)
    _add_key_value_table(
        doc,
        [
            ("Autores", _stringify(identification.get("authors") or "No detectado")),
            ("Ano", _stringify(identification.get("year") or "No detectado")),
            ("Tipo", _stringify(identification.get("document_type") or "No detectado")),
        ],
    )

    doc.add_heading("Investigacion", level=3)
    _add_key_value_table(
        doc,
        [
            ("Objetivo", _stringify(research.get("main_objective") or "No se menciona")),
            ("Pregunta", _stringify(research.get("research_question") or "No se menciona")),
            ("Problema", _stringify(research.get("problem") or "No se menciona")),
        ],
    )

    doc.add_heading("Metodologia", level=3)
    _add_key_value_table(
        doc,
        [
            ("Enfoque", _stringify(methodology.get("approach") or "No se menciona")),
            ("Diseno", _stringify(methodology.get("design") or "No se menciona")),
            ("Poblacion", _stringify(methodology.get("population") or "No se menciona")),
            ("Muestra", _stringify(methodology.get("sample") or "No se menciona")),
        ],
    )
    _add_bullets(doc, "Tecnicas", methodology.get("techniques"))
    _add_bullets(doc, "Criterios de seleccion", methodology.get("selection_criteria"))
    _render_evidence(doc, methodology.get("evidence"))

    _add_bullets(doc, "Resultados", result.get("results"))
    _render_document_elements(doc, result)

    doc.add_heading("Conclusiones", level=3)
    _add_callout(doc, "Conclusion general", conclusions.get("general_conclusion"))
    _add_label_value(doc, "Respuesta de investigacion", conclusions.get("research_answer"))
    _add_label_value(doc, "Implicancias", conclusions.get("implications"))
    _add_bullets(doc, "Limitaciones", conclusions.get("limitations"))
    _add_bullets(doc, "Recomendaciones futuras", conclusions.get("future_recommendations"))

    _render_quotes(doc, result.get("useful_quotes"))
    _add_callout(doc, "Aporte al campo", result.get("field_contribution"))
    doc.add_heading("Relevancia para tu tesis", level=3)
    _add_key_value_table(
        doc,
        [
            ("Nivel", _stringify(thesis.get("level") or "No detectado")),
            ("Motivo", _stringify(thesis.get("reason") or "No se menciona")),
            ("Uso posible", _stringify(thesis.get("possible_use") or "No se menciona")),
        ],
    )
    _add_bullets(doc, "Palabras clave", result.get("keywords"))


def _render_evidence(doc: Document, evidence: Any) -> None:
    if not isinstance(evidence, list) or not evidence:
        return
    doc.add_heading("Evidencia metodologica", level=3)
    rows = [["Campo", "Fragmento", "Fuente"]]
    for item in evidence:
        if isinstance(item, dict):
            rows.append([
                _stringify(item.get("field") or "Evidencia"),
                _stringify(item.get("excerpt") or ""),
                _stringify(item.get("source") or ""),
            ])
    _add_data_table(doc, rows)


def _render_quotes(doc: Document, quotes: Any) -> None:
    if not isinstance(quotes, list) or not quotes:
        return
    doc.add_heading("Citas utiles", level=3)
    rows = [["Cita", "Ubicacion", "Utilidad"]]
    for quote in quotes:
        if isinstance(quote, dict):
            rows.append([
                f'"{_stringify(quote.get("text") or "")}"',
                _stringify(quote.get("location") or ""),
                _stringify(quote.get("usefulness") or ""),
            ])
    _add_data_table(doc, rows)


def _render_document_elements(doc: Document, result: dict[str, Any]) -> None:
    elements = result.get("document_elements")
    if not isinstance(elements, list) or not elements:
        return
    doc.add_heading("Elementos detectados", level=3)
    rows = [["Tipo", "Descripcion", "Fuente"]]
    for element in elements:
        if isinstance(element, dict):
            rows.append([
                _stringify(element.get("type") or "Elemento"),
                _stringify(element.get("description") or ""),
                _stringify(element.get("source") or ""),
            ])
    _add_data_table(doc, rows)


def _add_data_table(doc: Document, rows: list[list[str]]) -> None:
    if len(rows) < 2:
        return
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for index, header in enumerate(rows[0]):
        cell = table.rows[0].cells[index]
        _set_cell_text(cell, header, bold=True, color=_WHITE)
        _set_cell_bg(cell, _HEADER_BG)

    for row_index, row_data in enumerate(rows[1:]):
        row = table.add_row()
        for col_index, text in enumerate(row_data):
            cell = row.cells[col_index]
            _set_cell_text(cell, text, size=8)
            if row_index % 2 == 1:
                _set_cell_bg(cell, _ALT_ROW_BG)

    doc.add_paragraph()
